"""Extract plain text from uploaded files (PDF, DOCX, TXT, CSV, MD, JSON)."""

from __future__ import annotations

import io
import json
import logging
from pathlib import Path

logger = logging.getLogger("file_extractor")

# Max characters extracted from any file to avoid overwhelming the LLM context
MAX_CHARS = 40_000

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".json"}


def supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


async def extract_text(filename: str, content: bytes) -> str:
    """Return plain text extracted from *content* (raw bytes of *filename*)."""
    ext = Path(filename).suffix.lower()

    if ext in (".txt", ".md"):
        text = content.decode("utf-8", errors="replace")

    elif ext == ".csv":
        text = _extract_csv(content)

    elif ext == ".json":
        text = _extract_json(content)

    elif ext == ".pdf":
        text = _extract_pdf(content)

    elif ext == ".docx":
        text = _extract_docx(content)

    else:
        raise ValueError(
            f"Unsupported file type '{ext}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    text = text.strip()
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + f"\n\n[... truncated at {MAX_CHARS} characters ...]"

    logger.info("Extracted %d chars from %s", len(text), filename)
    return text


# ─── Format-specific helpers ──────────────────────────────────────────────────

def _extract_csv(content: bytes) -> str:
    import csv

    lines = content.decode("utf-8", errors="replace").splitlines()
    reader = csv.reader(lines)
    rows = ["\t".join(row) for row in reader]
    return "\n".join(rows)


def _extract_json(content: bytes) -> str:
    try:
        data = json.loads(content.decode("utf-8", errors="replace"))
        return json.dumps(data, indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError("pypdf is required for PDF extraction. Run: pip install pypdf")

    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text)

    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise ImportError(
            "python-docx is required for Word document extraction. Run: pip install python-docx"
        )

    doc = docx.Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            cell_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cell_texts:
                paragraphs.append(" | ".join(cell_texts))

    return "\n".join(paragraphs)
